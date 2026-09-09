import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
from inventory_import import parse_batch_file, parse_inventory_rows


class InventoryImportTests(unittest.TestCase):
    def test_individual_plants_grouped_with_canonical_names(self):
        rows = [('Tag', 'Strain', 'Harvested'), ('a', ' chow mein ', 0),
                ('b', 'Chow Mein', 0), ('c', 'New Strain', 0)]
        self.assertEqual(parse_inventory_rows(rows, ['Chow Mein']), [('Chow Mein', 2), ('New Strain', 1)])

    def test_count_column(self):
        self.assertEqual(parse_inventory_rows([
            ['Inventory report'], ['Strain Name', 'Plant Count'],
            ['A', 12.0], ['a', '3'], ['B', 2],
        ]), [('A', 15), ('B', 2)])

    def test_invalid_inventory(self):
        for rows in (
            [['Tag'], ['a']], [['Strain']],
            [['Tag', 'Strain'], ['a', 'A'], ['a', 'A']],
            [['Tag', 'Strain'], ['', 'A']],
            [['Strain', 'Count'], ['A', 1.5]],
            [['Strain', 'Count'], ['A', 0]],
            [['Strain', 'Count'], ['A', None]],
            [['Strain', 'Count'], ['', 1]],
        ):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                parse_inventory_rows(rows)

    def test_metrc_incorrect_dimensions(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'inventory.XLSX'
            workbook = Workbook()
            workbook.active.append(['Tag', 'Strain', 'Location'])
            workbook.active.append(['a', 'A', 'Veg'])
            workbook.active.append(['b', 'A', 'Veg'])
            workbook.save(path)
            workbook.close()
            with ZipFile(path) as archive:
                contents = {name: archive.read(name) for name in archive.namelist()}
            contents['xl/worksheets/sheet1.xml'] = contents['xl/worksheets/sheet1.xml'].replace(b'A1:C3', b'A1:A3')
            with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
                for name, data in contents.items():
                    archive.writestr(name, data)
            self.assertEqual(parse_batch_file(path), [('A', 2)])

    def test_multiple_sheets_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'inventory.xlsx'
            workbook = Workbook()
            for sheet in (workbook.active, workbook.create_sheet()):
                sheet.append(['Strain'])
                sheet.append(['A'])
            workbook.save(path)
            workbook.close()
            with self.assertRaisesRegex(ValueError, 'Multiple inventory sheets'):
                parse_batch_file(path)

    def test_csv_still_supported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'batches.csv'
            path.write_text('strain,count\nA,12\n')
            self.assertEqual(parse_batch_file(path, ['A']), [('A', 12)])
